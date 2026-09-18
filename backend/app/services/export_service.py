import csv
import io
import re
from openpyxl import Workbook
from typing import Dict, Any, Tuple
from app.models.invoice import Invoice
from app.core.exceptions import InvoiceExportError

class ExportService:
    @staticmethod
    def _ensure_exportable(invoice: Invoice) -> Dict[str, Any]:
        status = invoice.status.value if hasattr(invoice.status, "value") else invoice.status
        if status != "EXTRACTED":
            raise InvoiceExportError(
                "Export is available after invoice extraction completes."
            )
        data = invoice.extracted_data
        if not isinstance(data, dict):
            raise InvoiceExportError("This invoice has no extracted data to export.")
        return data

    @staticmethod
    def _export_filename(invoice: Invoice, extension: str) -> str:
        stem = re.sub(r"[^A-Za-z0-9._-]+", "_", invoice.filename.rsplit(".", 1)[0])
        return f"{stem or 'invoice'}_export.{extension}"

    @staticmethod
    def _flatten_invoice_data(invoice: Invoice) -> Dict[str, Any]:
        data = ExportService._ensure_exportable(invoice)
            
        flat = {
            "Invoice ID": invoice.id,
            "Filename": invoice.filename,
            "Status": invoice.status.value if hasattr(invoice.status, 'value') else invoice.status,
            "Created At": invoice.created_at.isoformat() if invoice.created_at else "",
            "Invoice Number": data.get("invoice_number", ""),
            "Vendor Name": data.get("vendor_name", ""),
            "Vendor Address": data.get("vendor_address", ""),
            "Customer Name": data.get("customer_name", ""),
            "Invoice Date": data.get("invoice_date", ""),
            "Due Date": data.get("due_date", ""),
            "Currency": data.get("currency", ""),
            "Subtotal": data.get("subtotal", ""),
            "Tax": data.get("tax", ""),
            "Discount": data.get("discount", ""),
            "Total": data.get("total", ""),
            "Payment Terms": data.get("payment_terms", ""),
            "AI Confidence": data.get("confidence_score", invoice.overall_confidence or ""),
        }
        return flat
        
    @staticmethod
    def generate_csv(invoice: Invoice) -> Tuple[str, bytes]:
        """Returns filename and bytes for CSV export"""
        flat_data = ExportService._flatten_invoice_data(invoice)
        output = io.StringIO(newline="")
        writer = csv.DictWriter(output, fieldnames=list(flat_data.keys()))
        writer.writeheader()
        writer.writerow(flat_data)
        
        # Add line items
        data = invoice.extracted_data or {}
        items = data.get("invoice_items", []) if isinstance(data, dict) else []
        if items:
            output.write("\n")
            raw_writer = csv.writer(output)
            raw_writer.writerow(["Line Items"])
            raw_writer.writerow(["Description", "Quantity", "Unit Price", "Total"])
            
            for item in items:
                raw_writer.writerow([
                    item.get("description", ""),
                    item.get("quantity", ""),
                    item.get("unit_price", ""),
                    item.get("total", "")
                ])
        
        return ExportService._export_filename(invoice, "csv"), output.getvalue().encode("utf-8-sig")

    @staticmethod
    def generate_excel(invoice: Invoice) -> Tuple[str, bytes]:
        """Returns filename and bytes for Excel export"""
        flat_data = ExportService._flatten_invoice_data(invoice)
        wb = Workbook()
        ws_main = wb.active
        ws_main.title = "Invoice Summary"
        
        # Header for main sheet
        for col_idx, column_title in enumerate(flat_data.keys(), 1):
            ws_main.cell(row=1, column=col_idx, value=column_title)
        
        # Data for main sheet
        for col_idx, value in enumerate(flat_data.values(), 1):
            ws_main.cell(row=2, column=col_idx, value=value)
            
        # Line Items sheet
        data = ExportService._ensure_exportable(invoice)
        items = data.get("invoice_items", [])
        if items:
            ws_items = wb.create_sheet(title="Line Items")
            headers = ["Description", "Quantity", "Unit Price", "Total"]
            for col_idx, header in enumerate(headers, 1):
                ws_items.cell(row=1, column=col_idx, value=header)
                
            for row_idx, item in enumerate(items, 2):
                ws_items.cell(row=row_idx, column=1, value=item.get("description", ""))
                ws_items.cell(row=row_idx, column=2, value=item.get("quantity", ""))
                ws_items.cell(row=row_idx, column=3, value=item.get("unit_price", ""))
                ws_items.cell(row=row_idx, column=4, value=item.get("total", ""))
                
        output = io.BytesIO()
        wb.save(output)
        return ExportService._export_filename(invoice, "xlsx"), output.getvalue()
